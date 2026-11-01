"""Concours en plusieurs étapes : les admis au QCM passent à l'écrit, les
admis à l'écrit passent à l'oral.
"""
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import logging

import django

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "quizscan.settings")
os.environ["QUIZSCAN_TEST"] = "1"
django.setup()
from django.core.management import call_command  # noqa: E402
call_command("migrate", verbosity=0)

from django.contrib.auth.models import User
from django.test import Client
from django.test.utils import setup_test_environment
from django.urls import reverse

from grader import services
from grader.models import (Answer, ClassGroup, Question, Quiz, ScanBatch, SheetScan,
                           Student)
from test_commun import enseignant_complet

# nom : bonnes réponses au QCM (sur 10)
CANDIDATS = {"AMRI": 9, "BAKIR": 8, "CHERIF": 7, "DRISS": 5, "ESSID": 4, "FERCHICHI": 2}


def _noter(quiz, notes):
    """Copies déjà corrigées : {étudiant: note} sur la question « Note »
    (écrit / oral) ou {étudiant: nb de bonnes réponses} (QCM)."""
    lot = ScanBatch.objects.create(quiz=quiz, status="done", total_pages=len(notes))
    for s, note in notes.items():
        copie = SheetScan.objects.create(batch=lot, status="ok", student=s,
                                         source_name=f"{s.last_name}.jpg")
        if quiz.est_qcm:
            for n, q in enumerate(quiz.questions.order_by("order"), start=1):
                a = Answer(sheet=copie, question=q, detected_choice=0 if n <= note else 1)
                a.points_awarded = services._auto_points(q, a)
                a.save()
        else:
            Answer.objects.create(sheet=copie, question=services.question_note(quiz),
                                  points_awarded=note, manually_set=True)


def main():
    logging.disable(logging.WARNING)
    setup_test_environment()
    Quiz.objects.filter(owner__username="prof_etapes").delete()
    ClassGroup.objects.filter(owner__username="prof_etapes").delete()
    User.objects.filter(username="prof_etapes").delete()
    prof = enseignant_complet(User.objects.create_user("prof_etapes", password="x-pass-123"))
    c = Client()
    c.force_login(prof)

    # ---- étape 1 : QCM du concours, note minimale 6 / 10 -> 3 admis
    groupe = ClassGroup.objects.create(name="Concours — Étapes QCM", owner=prof)
    qcm = Quiz.objects.create(title="Étapes QCM", class_group=groupe, owner=prof,
                              auto_enroll=True, id_mode="sticker", note_admission=6)
    for i in range(1, 11):
        Question.objects.create(quiz=qcm, order=i, qtype="qcm", correct_choice=0)
    cands = {nom: Student.objects.create(class_group=groupe, last_name=nom,
                                         first_name="X", student_number=f"55{i:02d}")
             for i, nom in enumerate(CANDIDATS)}
    _noter(qcm, {cands[n]: v for n, v in CANDIDATS.items()})
    assert [s.last_name for s in services.admis_de(qcm)] == ["AMRI", "BAKIR", "CHERIF"]
    page = c.get(reverse("quiz_results", args=[qcm.pk])).content.decode()
    assert "Créer l'écrit des 3 admis" in page and "Créer l'oral des 3 admis" in page
    print("  QCM : 3 admis ; « Créer l'écrit des 3 admis » sur la page des résultats OK")

    # ---- étape 2 : l'écrit, créé depuis ce bouton (formulaire prérempli)
    url = reverse("quiz_create") + f"?concours=1&precedente={qcm.pk}&type=ecrit"
    form = c.get(url).context["form"]
    assert form.initial["epreuve_precedente"] == qcm.pk
    assert form.initial["type_epreuve"] == "ecrit"
    r = c.post(url, {"title": "Étapes écrit", "type_epreuve": "ecrit",
                     "epreuve_precedente": qcm.pk, "language": "fr", "note_sur": "20",
                     "sheet_mode": "grid", "id_mode": "sticker", "id_digits": "6",
                     "grading_mode": "immediate", "wrong_penalty": "0"}, follow=True)
    ecrit = Quiz.objects.get(title="Étapes écrit")
    assert ecrit.epreuve_precedente_id == qcm.pk and ecrit.class_group_id != groupe.pk
    inscrits = sorted(ecrit.class_group.students.values_list("last_name", flat=True))
    assert inscrits == ["AMRI", "BAKIR", "CHERIF"], inscrits
    assert ecrit.class_group.students.get(last_name="AMRI").student_number == "5500"
    assert "3 admis de « Étapes QCM » ajoutés" in r.content.decode()
    page = c.get(reverse("quiz_detail", args=[ecrit.pk])).content.decode()
    # liste à jour : pas d'encart ni de bouton, l'en-tête rappelle l'origine
    assert "candidats : admis de" in page and "Mettre à jour la liste" not in page
    print("  écrit : liste des candidats = les 3 admis du QCM OK")

    # ---- étape 3 : l'oral (épreuve de classe), admis de l'écrit (≥ 10 / 20)
    e = {s.last_name: s for s in ecrit.class_group.students.all()}
    ecrit.note_admission = 10
    ecrit.save(update_fields=["note_admission"])
    _noter(ecrit, {e["AMRI"]: 14.5, e["BAKIR"]: 8.0, e["CHERIF"]: 12.25})
    r = c.post(reverse("quiz_create") + f"?precedente={ecrit.pk}&type=oral",
               {"title": "Étapes oral", "type_epreuve": "oral",
                "epreuve_precedente": ecrit.pk, "language": "fr", "note_sur": "20",
                "sheet_mode": "grid", "id_mode": "grid", "id_digits": "6",
                "wrong_penalty": "0"})
    assert r.status_code == 302, r.context["form"].errors
    oral = Quiz.objects.get(title="Étapes oral")
    inscrits = sorted(oral.class_group.students.values_list("last_name", flat=True))
    assert inscrits == ["AMRI", "CHERIF"], inscrits
    page = c.get(reverse("quiz_detail", args=[oral.pk])).content.decode()
    assert "AMRI X" in page and "BAKIR" not in page.split("Notes de l'oral")[1]
    print("  oral : liste des candidats = les 2 admis de l'écrit OK")

    # ---- la note minimale du QCM baisse : un nouvel admis à faire passer
    qcm.note_admission = 5
    qcm.save(update_fields=["note_admission"])
    page = c.get(reverse("quiz_detail", args=[ecrit.pk])).content.decode()
    assert "1 nouvel admis" in page, "nouvel admis non signalé"
    c.post(reverse("quiz_detail", args=[ecrit.pk]), {"passer_admis": "1"})
    assert ecrit.class_group.students.filter(last_name="DRISS").exists()
    c.post(reverse("quiz_detail", args=[ecrit.pk]), {"passer_admis": "1"})
    assert ecrit.class_group.students.count() == 4, "doublon"
    page = c.get(reverse("quiz_detail", args=[ecrit.pk])).content.decode()
    assert "Mettre à jour la liste" not in page and "nouvel admis" not in page
    print("  note minimale baissée : nouvel admis signalé, ajouté, sans doublon OK")

    # ---- résultats de tout le concours, depuis n'importe quelle étape
    oral.note_admission = 10
    oral.save(update_fields=["note_admission"])
    o = {s.last_name: s for s in oral.class_group.students.all()}
    _noter(oral, {o["AMRI"]: 15, o["CHERIF"]: 9})
    attendu = {"AMRI": "Admis au concours", "CHERIF": "Non admis : Étapes oral",
               "BAKIR": "Non admis : Étapes écrit", "DRISS": "Non noté : Étapes écrit",
               "ESSID": "Non admis : Étapes QCM", "FERCHICHI": "Non admis : Étapes QCM"}
    for depart in (qcm, ecrit, oral):
        r = c.get(reverse("quiz_parcours", args=[depart.pk]))
        assert r.status_code == 200
        p = r.context["parcours"]
        assert [q.pk for q in p["epreuves"]] == [qcm.pk, ecrit.pk, oral.pk]
        finals = {x["etudiant"].last_name: x["final"] for x in p["candidats"]}
        assert finals == attendu, finals
        assert p["n_admis_final"] == 1
    ordre = [x["etudiant"].last_name for x in p["candidats"]]
    assert ordre == ["AMRI", "CHERIF", "BAKIR", "DRISS", "ESSID", "FERCHICHI"], ordre
    page = r.content.decode()
    assert "Admis au concours" in page and "non convoqué" in page
    assert [x["etat"] for x in p["candidats"][-1]["cases"]] == ["note", "non_convoque", "non_convoque"]
    # la recherche filtre, l'export Excel reprend tout
    r = c.get(reverse("quiz_parcours", args=[ecrit.pk]) + "?q=driss")
    assert [x["etudiant"].last_name for x in r.context["candidats_page"]] == ["DRISS"]
    r = c.get(reverse("quiz_parcours", args=[ecrit.pk]) + "?format=xlsx")
    assert r["Content-Type"].startswith("application/vnd.openxmlformats")
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(row=2, column=ws.max_column).value == "Résultat final"
    assert ws.cell(row=3, column=1).value == "AMRI"
    assert ws.cell(row=3, column=ws.max_column).value == "Admis au concours"
    assert ws.max_row == 2 + len(attendu)
    # la page des résultats d'une étape mène aux résultats du concours
    page = c.get(reverse("quiz_results", args=[ecrit.pk])).content.decode()
    assert "Voir les résultats de tout le concours" in page
    # retour vers la liste des résultats ; l'épreuve reste à un clic (titre)
    assert f'href="{reverse("resultats_list")}">← Retour aux résultats' in page
    assert f'href="{reverse("quiz_detail", args=[ecrit.pk])}" title="Ouvrir l' in page
    assert reverse("quiz_parcours", args=[ecrit.pk]) in page
    for q in (ecrit, oral):
        page = c.get(reverse("quiz_detail", args=[q.pk])).content.decode()
        assert reverse("quiz_parcours", args=[q.pk]) in page
    seul = Quiz.objects.create(title="Hors concours", class_group=groupe, owner=prof)
    page = c.get(reverse("quiz_results", args=[seul.pk])).content.decode()
    assert "Résultats du concours" not in page
    print("  résultats du concours : parcours de chaque candidat, résultat final, Excel OK")

    # une épreuve d'un collègue n'est pas proposée comme épreuve précédente
    autre = enseignant_complet(User.objects.get_or_create(username="prof_etapes2")[0])
    c2 = Client()
    c2.force_login(autre)
    form = c2.get(reverse("quiz_create")).context["form"]
    assert qcm not in form.fields["epreuve_precedente"].queryset
    assert c2.get(reverse("quiz_parcours", args=[ecrit.pk])).status_code == 404
    print("  épreuve d'un collègue : jamais proposée OK")
    print("\n✅ TEST DES CONCOURS EN PLUSIEURS ÉTAPES RÉUSSI")


if __name__ == "__main__":
    main()
