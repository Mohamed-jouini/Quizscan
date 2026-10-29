"""Une QCM = un point ; le résultat se lit en bonnes réponses, en pourcentage,
et admis / non admis — sur la page des résultats, l'espace public, les
bulletins, la liste des admis et l'export Excel.
"""
import io
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import django

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "quizscan.settings")
os.environ["QUIZSCAN_TEST"] = "1"   # base de test séparée (data/test_db.sqlite3)
django.setup()
from django.core.management import call_command  # noqa: E402
call_command("migrate", verbosity=0)

import importlib

import pymupdf
from django.apps import apps
from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from grader import services
from grader.models import (Answer, ClassGroup, Question, Quiz, ScanBatch, SheetScan,
                           Student)
from test_commun import enseignant_complet

# nom : bonnes réponses sur 6
CANDIDATS = {"AMRI": 5, "BAKIR": 3, "CHERIF": 2}


def _epreuve(prof, avec_redaction=False):
    groupe = ClassGroup.objects.create(name=f"UNPT{avec_redaction:d}", owner=prof)
    quiz = Quiz.objects.create(title=f"Un point {avec_redaction:d}", class_group=groupe,
                               owner=prof, sheet_mode="grid", id_mode="grid",
                               note_admission=3)
    qcm = [Question.objects.create(quiz=quiz, order=n, qtype="qcm", points=3,
                                   num_choices=4, correct_choice=0)
           for n in range(1, 7)]
    if avec_redaction:
        Question.objects.create(quiz=quiz, order=7, qtype="open", points=4)
    lot = ScanBatch.objects.create(quiz=quiz, status="done", total_pages=3)
    for i, (nom, justes) in enumerate(CANDIDATS.items()):
        s = Student.objects.create(class_group=groupe, last_name=nom, first_name="X",
                                   student_number=f"88{avec_redaction:d}{i}")
        copie = SheetScan.objects.create(batch=lot, status="ok", student=s,
                                         source_name=f"{nom}.jpg")
        for n, q in enumerate(qcm, start=1):
            a = Answer(sheet=copie, question=q, detected_choice=0 if n <= justes else 1)
            a.points_awarded = services._auto_points(q, a)
            a.save()
    return quiz


def main():
    Quiz.objects.filter(owner__username="prof_unpt").delete()
    ClassGroup.objects.filter(owner__username="prof_unpt").delete()
    User.objects.filter(username="prof_unpt").delete()
    prof = enseignant_complet(User.objects.create_user("prof_unpt", password="x-pass-123"))
    c = Client()
    c.force_login(prof)

    quiz = _epreuve(prof)
    assert set(quiz.questions.values_list("points", flat=True)) == {1.0}, \
        "une QCM créée avec 3 points doit valoir 1 point"
    rows = {r["student"].last_name: r for r in services.compute_results(quiz)}
    a = rows["AMRI"]
    assert (a["correct"], a["n_qcm"], a["pct"], a["score"]) == (5, 6, 83, 5.0), a
    assert not quiz.note_distincte
    print("  une QCM vaut 1 point : 5 bonnes réponses = 5 / 6 = 83 % OK")

    page = c.get(reverse("quiz_results", args=[quiz.pk])).content.decode()
    for attendu in ("<th>Bonnes réponses</th>", "<th>Pourcentage</th>", "5 / 6",
                    "83 %", "50 %", "Réussi", "Non réussi",
                    "au moins 3 bonnes réponses sur 6"):
        assert attendu in page, f"« {attendu} » absent de la page des résultats"
    assert "<th>Note</th>" not in page, "la note répète les bonnes réponses"
    print("  page des résultats : bonnes réponses, pourcentage, réussi / non réussi OK")

    xl = load_workbook(io.BytesIO(c.get(reverse("quiz_results_xlsx", args=[quiz.pk])).content))
    ws = xl["Résultats"]
    entetes = [x.value for x in ws[2]]
    ligne = {r[0].value: r for r in ws.iter_rows(min_row=3)}["AMRI"]
    assert round(ligne[entetes.index("Pourcentage")].value, 2) == 0.83
    assert ligne[entetes.index("Correctes")].value == 5
    assert ligne[entetes.index("Résultat")].value == "Admis"
    print("  Excel : colonnes « Correctes », « Pourcentage », « Résultat » OK")

    pdf = c.get(reverse("quiz_bulletins_pdf", args=[quiz.pk]) + "?student="
                + str(a["student"].pk)).content
    texte = pymupdf.open(stream=pdf, filetype="pdf")[0].get_text()
    assert "BONNES RÉPONSES" in texte and "5 / 6" in texte and "83 %" in texte
    texte = pymupdf.open(stream=c.get(reverse("quiz_admis_pdf", args=[quiz.pk])).content,
                         filetype="pdf")[0].get_text()
    assert "Bonnes rép." in texte and "5 / 6" in texte and "83 %" in texte
    assert "CHERIF" not in texte
    print("  bulletin et liste des réussis en PDF : bonnes réponses et pourcentage OK")

    c.post(reverse("quiz_publication", args=[quiz.pk]), {"resultats_publies": "1",
                                                         "admis_publies": "1"})
    cache.clear()
    anonyme = Client()
    page = anonyme.get(reverse("public_resultat") + "?numero=8800").content.decode()
    assert "<b>5</b> / 6 bonnes réponses" in page and "83 %" in page and "Réussi" in page
    page = anonyme.get(reverse("public_resultat") + "?numero=8802").content.decode()
    assert "<b>2</b> / 6 bonnes réponses" in page and "33 %" in page and "Non réussi" in page
    page = anonyme.get(reverse("public_epreuve", args=[quiz.pk])).content.decode()
    assert "<b>5</b> / 6" in page and "83 %" in page
    cache.clear()
    print("  espace public : bonnes réponses, pourcentage, réussi / non réussi OK")

    # avec une réponse manuscrite, la note ne se confond plus avec les bonnes réponses
    quiz2 = _epreuve(prof, avec_redaction=True)
    assert quiz2.questions.get(order=7).points == 4.0
    assert quiz2.note_distincte
    page = c.get(reverse("quiz_results", args=[quiz2.pk])).content.decode()
    assert "<th>Note</th>" in page and "/ 10" in page
    print("  réponse manuscrite : barème libre, colonne « Note » en plus OK")

    # reprise des épreuves existantes (migration 0028)
    q = quiz.questions.get(order=1)
    Question.objects.filter(pk=q.pk).update(points=3)
    Answer.objects.filter(question=q, detected_choice=0).update(points_awarded=3)
    migration = importlib.import_module("grader.migrations.0028_qcm_un_point")
    migration.un_point(apps, None)
    q.refresh_from_db()
    assert q.points == 1.0
    assert set(Answer.objects.filter(question=q, detected_choice=0)
               .values_list("points_awarded", flat=True)) == {1.0}
    print("  migration : QCM existantes ramenées à 1 point, notes recalculées OK")
    print("\n✅ TEST « UNE QCM = UN POINT » RÉUSSI")


if __name__ == "__main__":
    main()
