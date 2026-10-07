"""Export Excel des résultats."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def results_workbook(quiz, rows, stats=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Résultats"

    headers = ["Nom", "Prénom", "N° inscription", "Note", "Sur",
               "Répondues", "Correctes", "Fausses", "Sans réponse",
               "Manuscrites à corriger"]
    ws.append([f"{quiz.title} — {quiz.class_group.name}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="1F4E79")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=2, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        s = r["student"]
        ws.append([s.last_name, s.first_name, s.student_number,
                   r["score"], r["max_score"], r["answered"], r["correct"],
                   r["wrong"], r["blank"], r["pending_open"]])

    widths = [22, 18, 14, 8, 6, 11, 10, 8, 13, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"

    if stats:
        ws2 = wb.create_sheet("Statistiques")
        ws2.append(["Étudiants", stats["n"]])
        ws2.append(["Moyenne", stats["mean"]])
        ws2.append(["Médiane", stats["median"]])
        ws2.append(["Écart-type", stats["stdev"]])
        ws2.append(["Meilleure note", stats["best"]])
        ws2.append(["Note la plus basse", stats["worst"]])
        ws2.append(["Barème total", stats["max_score"]])
        if stats["suspectes"]:
            ws2.append(["Bonne réponse à vérifier (questions)",
                        ", ".join(str(n) for n in stats["suspectes"])])
        ws2.append([])

        entetes = ["Question", "Type", "Réussite %", "Justes", "Fausses",
                   "Sans réponse", "Copies notées", "Discrimination",
                   "Lecture", "Répartition des réponses"]
        ws2.append(entetes)
        # Position calculée et non écrite en dur : le résumé ci-dessus a déjà
        # gagné deux lignes, et en gagnera peut-être d'autres.
        ligne_entete = ws2.max_row
        for c in range(1, len(entetes) + 1):
            cell = ws2.cell(row=ligne_entete, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")

        for s in stats["questions"]:
            q = s["question"]
            indice = s.get("discrimination")
            lecture = s.get("discrimination_label", "")
            if s.get("corrige_suspect"):
                lecture = "suspecte — vérifier la bonne réponse"
            repartition = " · ".join(
                f"{c['lettre']} {c['n']}" for c in s.get("repartition", []))
            if q.qtype == "qcm":
                # une question manuscrite n'a ni juste, ni faux, ni vide :
                # on laisse les colonnes à blanc au lieu d'afficher des 0
                ws2.append([q.order, "QCM", s["success"], s["correct"],
                            s["wrong"], s["blank"], s["n"],
                            indice, lecture, repartition])
            else:
                ws2.append([q.order, "Manuscrite", s["success"], None, None,
                            None, s["graded"], None, "", ""])
        for i, w in enumerate([18, 12, 12, 9, 9, 13, 14, 14, 34, 30], start=1):
            ws2.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
